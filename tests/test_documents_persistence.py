"""Phase 5 scoped ingestion, located citations, memory and real process restart gates."""

import asyncio
import hashlib
import json
import os
import re
import sqlite3
import subprocess
import sys
from tempfile import TemporaryFile

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from evidenceforge.models.research import FixtureResearchModel
from evidenceforge.reporting import render_report
from evidenceforge.retrieval.contracts import RetrievalBatch
from evidenceforge.retrieval.fixtures import FixtureRetrievalAdapter
from evidenceforge.schemas import ResearchRequest, SourceType
from evidenceforge.services.research import ResearchService
from evidenceforge.services.retrieval import RetrievalService
from evidenceforge.storage.checkpoints import CheckpointFailure, CheckpointRepository
from evidenceforge.storage.documents import DocumentFailure, DocumentOptions, DocumentStore
from evidenceforge.storage.reports import ReportRepository

QUESTION = "Compare solar and wind electricity output."
SOLAR = "Synthetic solar electricity output depends on sunlight."
WIND = "Synthetic wind electricity output depends on wind speed."


def make_pdf(path, *, blank=False):
    writer = PdfWriter()
    for text in [] if blank else [SOLAR, WIND]:
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
        )
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 50 740 Td ({text}) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
    writer.add_blank_page(width=612, height=792)
    writer.write(path)


def make_doc(path, kind):
    if kind == "txt":
        path.write_text(SOLAR + "\n\n" + WIND, encoding="utf-8")
    elif kind == "pdf":
        make_pdf(path)
    else:
        document = Document()
        document.add_paragraph(SOLAR)
        document.add_paragraph(WIND)
        document.add_table(rows=1, cols=1).cell(
            0, 0
        ).text = "Synthetic table: solar electricity output is illustrative."
        document.save(path)


@pytest.mark.parametrize("kind", ["txt", "pdf", "docx"])
def test_ingestion_and_located_report(tmp_path, kind):
    path = tmp_path / f"synthetic.{kind}"
    make_doc(path, kind)
    store = DocumentStore(tmp_path / "data" / "documents")
    record = store.ingest(path, "owner_a")
    assert record.binary_hash == hashlib.sha256(path.read_bytes()).hexdigest()
    assert store.ingest(path, "owner_a") == record
    chunks = store.chunks("owner_a", record.id)
    assert all(
        chunk.content_hash == hashlib.sha256(chunk.text.encode()).hexdigest() for chunk in chunks
    )
    state = asyncio.run(
        ResearchService(tmp_path / "data").run(
            ResearchRequest(question=QUESTION, owner_scope="owner_a", document_ids=(record.id,))
        )
    )
    assert state["run_status"] == "completed"
    assert all(s["source_type"] == "document" for s in state["source_records"].values())
    assert all(s["owner_scope"] == "owner_a" for s in state["source_records"].values())
    assert all(
        e["locator"]["document_start_offset"] is not None for e in state["evidence_store"].values()
    )
    if kind == "pdf":
        assert {e["locator"]["page"] for e in state["evidence_store"].values()} == {1, 2}
        assert "page 2" in render_report(state)
        assert record.status == "partial"
    else:
        assert "paragraph 2" in render_report(state)
    assert "SHA-256" in render_report(state)
    assert "Offline fixture reasoning over local document" in render_report(state)


def test_textless_pdf_discloses_ocr(tmp_path):
    path = tmp_path / "blank.pdf"
    make_pdf(path, blank=True)
    store = DocumentStore(tmp_path / "documents")
    record = store.ingest(path)
    assert record.status == "textless"
    assert record.chunk_count == 0
    assert "OCR is unavailable" in record.warnings[0]
    assert not store.search("local", (record.id,), "electricity")


def test_owner_and_requested_document_filters(tmp_path):
    store = DocumentStore(tmp_path / "documents")
    a = tmp_path / "a.txt"
    a.write_text(SOLAR, encoding="utf-8")
    b = tmp_path / "b.txt"
    b.write_text("Private owner B secret electricity output.", encoding="utf-8")
    ra, rb = store.ingest(a, "owner_a"), store.ingest(b, "owner_b")
    with pytest.raises(DocumentFailure):
        store.get("owner_a", rb.id)
    with pytest.raises(DocumentFailure):
        store.search("owner_a", (rb.id,), "electricity")
    assert all(c.document_id == ra.id for c in store.search("owner_a", (ra.id,), "electricity"))
    assert not store.search("owner_a", (), "electricity")
    assert store.ingest(a, "owner_b").id != ra.id


@pytest.mark.parametrize(
    "suffix,data", [("exe", b"x"), ("pdf", b"bad pdf"), ("docx", b"bad zip"), ("txt", b"\xff\x00")]
)
def test_invalid_documents_rejected_without_registration(tmp_path, suffix, data):
    path = tmp_path / f"invalid.{suffix}"
    path.write_bytes(data)
    store = DocumentStore(tmp_path / "documents")
    with pytest.raises(DocumentFailure):
        store.ingest(path)
    with store.connect() as conn:
        assert conn.execute("SELECT COUNT(*) FROM documents").fetchone()[0] == 0


def test_limits_and_exact_overlapping_chunks(tmp_path):
    path = tmp_path / "long.txt"
    text = "solar electricity output " * 30
    path.write_text(text, encoding="utf-8")
    with pytest.raises(DocumentFailure):
        DocumentStore(tmp_path / "small", DocumentOptions(max_file_bytes=10)).ingest(path)
    store = DocumentStore(
        tmp_path / "documents", DocumentOptions(chunk_chars=100, overlap_chars=20)
    )
    record = store.ingest(path)
    chunks = sorted(store.chunks("local", record.id), key=lambda c: c.start_offset)
    assert chunks[1].start_offset == 80
    assert all(text[c.start_offset : c.end_offset] == c.text for c in chunks)
    assert store.search("local", (record.id,), "output", 2)


def test_corrupted_registered_document_is_not_retrieved(tmp_path):
    path = tmp_path / "data.txt"
    path.write_text(SOLAR, encoding="utf-8")
    store = DocumentStore(tmp_path / "documents")
    record = store.ingest(path)
    with store.connect() as conn:
        conn.execute("UPDATE documents SET text='tampered'")
    with pytest.raises(DocumentFailure):
        store.search("local", (record.id,), "solar")


def test_checkpoint_resume_does_not_repeat_completed_nodes(tmp_path):
    first_model = FixtureResearchModel()
    first_adapter = FixtureRetrievalAdapter()
    first = ResearchService(
        tmp_path,
        model=first_model,
        retrieval=RetrievalService(tmp_path, {SourceType.WEB: first_adapter}),
    )
    paused = asyncio.run(
        first.run(ResearchRequest(question=QUESTION), pause_after="source_processor")
    )
    assert paused["run_status"] == "running" and paused["iteration_count"] == 1
    assert not first.report_path(paused).exists()
    second_model, second_adapter = FixtureResearchModel(), FixtureRetrievalAdapter()
    second = ResearchService(
        tmp_path,
        model=second_model,
        retrieval=RetrievalService(tmp_path, {SourceType.WEB: second_adapter}),
    )
    final = asyncio.run(second.resume(paused["thread_id"]))
    assert final["run_status"] == "completed"
    assert final["run_id"] == paused["run_id"]
    assert second_adapter.calls == second_model.plan_calls == second_model.query_calls == 0
    assert final["iteration_count"] == 1
    assert set(final["calls"]) == set(paused["calls"])
    assert asyncio.run(second.resume(paused["thread_id"])) == final


def test_foreign_thread_and_configuration_mismatch_rejected(tmp_path):
    service = ResearchService(tmp_path)
    paused = asyncio.run(
        service.run(
            ResearchRequest(question=QUESTION, owner_scope="owner_a"), pause_after="query_generator"
        )
    )
    with pytest.raises(CheckpointFailure):
        asyncio.run(service.resume(paused["thread_id"], "owner_b"))
    from evidenceforge.research import ResearchOptions

    other = ResearchService(tmp_path, options=ResearchOptions(recursion_limit=32))
    with pytest.raises(CheckpointFailure):
        asyncio.run(other.resume(paused["thread_id"], "owner_a"))
    with pytest.raises(CheckpointFailure):
        asyncio.run(service.resume("thread_missing", "owner_a"))


def test_new_thread_resets_counters_and_terminal_data(tmp_path):
    service = ResearchService(tmp_path)
    final = asyncio.run(service.run(ResearchRequest(question=QUESTION)))
    fresh = asyncio.run(
        service.run(ResearchRequest(question=QUESTION), pause_after="query_generator")
    )
    assert fresh["thread_id"] != final["thread_id"] and fresh["run_id"] != final["run_id"]
    assert fresh["iteration_count"] == fresh["revision_count"] == 0
    assert fresh["draft_report"] is fresh["critic_verdict"] is fresh["final_report"] is None
    assert not fresh["evidence_store"] and not fresh["citation_audit"]
    assert any(c["kind"] == "prior_report" for c in fresh["context_refs"])


def test_reports_are_scoped_context_not_independent_evidence(tmp_path):
    service = ResearchService(tmp_path)
    state = asyncio.run(service.run(ResearchRequest(question=QUESTION, owner_scope="owner_a")))
    repository = ReportRepository(tmp_path / "memory")
    prior = repository.list("owner_a", query="solar")[0]
    assert prior.run_id == state["run_id"] and not prior.is_independent_evidence
    assert not repository.list("owner_b")
    with pytest.raises(DocumentFailure):
        service.documents.get("owner_a", prior.id)
    empty = ResearchService(
        tmp_path,
        retrieval=RetrievalService(
            tmp_path, {SourceType.WEB: FixtureRetrievalAdapter(RetrievalBatch(results=()))}
        ),
    )
    new = asyncio.run(empty.run(ResearchRequest(question=QUESTION, owner_scope="owner_a")))
    assert any(c["kind"] == "prior_report" for c in new["context_refs"])
    assert not new["evidence_store"] and new["run_status"] == "insufficient_evidence"


def cli(*args):
    env = dict(os.environ)
    env["EVIDENCEFORGE_MODEL_PROVIDER"] = "fixture"
    command = [sys.executable, "-m", "evidenceforge.research_demo", *map(str, args)]
    # The Windows venv launcher can leave a captured pipe open in its child.
    # A file preserves diagnostics without waiting on that pipe reader.
    with TemporaryFile(mode="w+t", encoding="utf-8") as output:
        try:
            result = subprocess.run(
                command, stdout=output, stderr=subprocess.STDOUT, env=env, timeout=180
            )
        except subprocess.TimeoutExpired as exc:
            output.seek(0)
            raise AssertionError(f"Research subprocess timed out: {output.read()}") from exc
        output.seek(0)
        rendered = output.read()
    assert result.returncode == 0, rendered
    return rendered


def test_actual_process_restart_resume(tmp_path):
    output = cli(
        QUESTION,
        "--data-dir",
        tmp_path,
        "--demo",
        "second-pass",
        "--pause-after",
        "source_processor",
    )
    thread = re.search(r"Thread: (thread_[a-f0-9]+)", output).group(1)
    result = cli("--data-dir", tmp_path, "--resume", thread)
    assert "Retrieval passes: 2" in result
    assert "Status: `completed`" in result
    paths = list((tmp_path / "reports").glob("*/state.json"))
    assert len(paths) == 1
    state = json.loads(paths[0].read_text(encoding="utf-8"))
    assert state["thread_id"] == thread
    with sqlite3.connect(CheckpointRepository(tmp_path / "checkpoints", "local").path) as conn:
        assert (
            conn.execute(
                "SELECT COUNT(*) FROM checkpoints WHERE thread_id=?", (thread,)
            ).fetchone()[0]
            > 1
        )


def test_active_thread_lock_rejects_second_writer(tmp_path):
    repository = CheckpointRepository(tmp_path / "checkpoints", "local")
    with repository.exclusive("thread_test"):
        with pytest.raises(CheckpointFailure):
            with repository.exclusive("thread_test"):
                pass
    with repository.exclusive("thread_test"):
        pass


def test_changed_chunk_locator_is_rejected(tmp_path):
    path = tmp_path / "data.txt"
    path.write_text(SOLAR, encoding="utf-8")
    store = DocumentStore(tmp_path / "documents")
    record = store.ingest(path)
    with store.connect() as conn:
        row = conn.execute("SELECT id,record FROM chunks").fetchone()
        chunk = json.loads(row[1])
        chunk["paragraph"] = 999
        conn.execute("UPDATE chunks SET record=? WHERE id=?", (json.dumps(chunk), row[0]))
    with pytest.raises(DocumentFailure):
        store.chunks("local", record.id)


def test_encrypted_pdf_is_rejected(tmp_path):
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.encrypt("synthetic_password")
    path = tmp_path / "encrypted.pdf"
    writer.write(path)
    with pytest.raises(DocumentFailure, match="Encrypted"):
        DocumentStore(tmp_path / "documents").ingest(path)


def test_document_checkpoint_survives_source_file_changes(tmp_path):
    path = tmp_path / "original.txt"
    make_doc(path, "txt")
    service = ResearchService(tmp_path / "data")
    record = service.documents.ingest(path, "owner_a")
    paused = asyncio.run(
        service.run(
            ResearchRequest(question=QUESTION, owner_scope="owner_a", document_ids=(record.id,)),
            pause_after="source_processor",
        )
    )
    path.write_text("The original file was replaced after retrieval.", encoding="utf-8")
    restarted = ResearchService(tmp_path / "data")
    final = asyncio.run(restarted.resume(paused["thread_id"], "owner_a"))
    assert final["run_status"] == "completed"
    assert set(final["calls"]) == set(paused["calls"])
    assert "sunlight" in render_report(final)
    assert all(s["document_hash"] == record.binary_hash for s in final["source_records"].values())


def test_thread_checkpoints_are_isolated(tmp_path):
    from langgraph.checkpoint.sqlite.aio import AsyncSqliteSaver

    service = ResearchService(tmp_path)
    a = asyncio.run(service.run(ResearchRequest(question=QUESTION), pause_after="query_generator"))
    b = asyncio.run(service.run(ResearchRequest(question=QUESTION), pause_after="query_generator"))
    final = asyncio.run(service.resume(a["thread_id"]))
    assert final["run_id"] == a["run_id"]

    async def read_b():
        repository = CheckpointRepository(tmp_path / "checkpoints", "local")
        async with AsyncSqliteSaver.from_conn_string(str(repository.path)) as saver:
            result = await saver.aget_tuple({"configurable": {"thread_id": b["thread_id"]}})
            return result.checkpoint["channel_values"]

    state = asyncio.run(read_b())
    assert state["iteration_count"] == 0
    assert state["run_id"] == b["run_id"]
    assert state["final_report"] is None
